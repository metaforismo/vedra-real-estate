"""Decision aggregation uses observed evidence, with no market-value assumptions."""
import io
from uuid import uuid4
from datetime import date,timedelta
from openpyxl import load_workbook
from app.db import now,dump
from app.services.asset_evidence import AssetEvidence
from app.services.market_references import MarketReferences
from app.services.exports import export_xlsx
from app.services.store import property_dict
from app.services.today import queue
from support.decision import decision_dataset


def records(api):
    app,c,s=api;db=app.state.db
    user=c.get('/api/auth/me').json()['user']['id']
    if db.one("SELECT id FROM sources WHERE id='qa-evidence-0'"):
        # One module fixture, each test explicitly resets only its own mutations.
        subject=db.one("SELECT id FROM properties WHERE listing_key='Asset da approfondire'")['id']
        linked=db.one("SELECT id FROM properties WHERE listing_key='Stesso asset, altra fonte'")['id']
        comps=[x['id'] for x in db.all("SELECT id FROM properties WHERE zone='QA · Zona confronto' AND id NOT IN (?,?)",(subject,linked))]
        return db,c,s,user,subject,linked,comps
    subject,linked,comps=decision_dataset(db,s,user)
    return db,c,s,user,subject,linked,comps


def test_cross_source_dossier_preserves_differences_and_exports(api):
    db,c,s,user,subject,linked,_=records(api)
    p=c.get('/api/properties/'+subject).json();cross=p['cross_sources']
    assert cross['count']==2 and cross['conflicts']==['Prezzi richiesti diversi']
    assert {row['price'] for row in cross['entries']}=={290000,310000}
    assert len({row['contact']['telephone'] for row in cross['entries']})==2
    wb=load_workbook(io.BytesIO(export_xlsx([p],db)),data_only=True)
    sheet=wb['Fonti dello stesso asset']
    assert sheet.max_row==3 and {sheet['E2'].value,sheet['E3'].value}=={290000,310000}
    selection=wb['Selezione'];headers=[cell.value for cell in selection[1]]
    assert 'Prezzi richiesti diversi' in selection.cell(2,headers.index('Motivi / verifiche')+1).value


def test_sample_diversity_quartiles_and_no_self_comparisons(api):
    db,c,s,user,subject,linked,_=records(api)
    groups=c.get('/api/properties/'+subject).json()['market_references']['groups']
    assert all(g['count']==3 and g['source_count']==3 for g in groups)
    first=groups[0]
    assert (first['median_sqm'],first['q1_sqm'],first['q3_sqm'])==(3400,3300,3500)
    assert first['oldest_observed']<=first['newest_observed']
    assert not any(row['id'] in (subject,linked) for g in groups for row in g['items'])


def test_today_deduplicates_confirmed_asset_and_checks_disputed_availability(api):
    db,c,s,user,subject,linked,_=records(api)
    ids={subject,linked}
    assert len([x for x in queue(db)['call'] if x['id'] in ids])==1
    db.execute("UPDATE properties SET availability='sold' WHERE id=?",(linked,))
    try:
        q=queue(db)
        assert not any(x['id'] in ids for x in q['call'])
        assert any(x['id']==subject and 'Disponibilità discordante' in x['checks'] for x in q['verify'])
    finally:db.execute("UPDATE properties SET availability='listed' WHERE id=?",(linked,))


def test_contact_followup_applies_to_confirmed_asset_across_sources(api):
    db,c,s,user,subject,linked,_=records(api)
    request=str(uuid4());future=(date.today()+timedelta(days=2)).isoformat()
    assert c.post('/api/properties/'+linked+'/contacts',json={'request_id':request,'outcome':'documents_requested','next_contact':future,'note':'QA documents'}).status_code==201
    try:
        q=queue(db)
        assert not any(x['id'] in (subject,linked) for x in q['call']+q['verify'])
    finally:db.execute('DELETE FROM contact_actions WHERE id=?',(request,))


def test_contradictory_identity_is_visible_not_silently_resolved(api):
    db,c,s,user,subject,linked,comps=records(api);third=comps[0]
    pairs=[(*sorted((linked,third)),'same_asset'),(*sorted((subject,third)),'distinct')]
    for a,b,decision in pairs:db.execute('INSERT INTO duplicate_reviews VALUES(?,?,?,?,?)',(a,b,decision,user,now()))
    try:
        evidence=AssetEvidence(db).for_property({'id':subject})
        assert evidence['identity_conflict'] and 'Collegamenti tra annunci da rivedere' in evidence['conflicts']
        assert not any(x['id'] in (subject,linked,third) for x in queue(db)['call'])
    finally:
        for a,b,_ in pairs:db.execute('DELETE FROM duplicate_reviews WHERE a=? AND b=?',(a,b))


def test_excel_history_keeps_observed_currency_and_breaks_incompatible_deltas(api):
    db,c,s,user,subject,linked,_=records(api)
    p=property_dict(db.one('SELECT * FROM properties WHERE id=?',(subject,)))
    p['observations']=[{'observed_at':now(),'price':price,'currency':currency,'transaction_type':transaction,'content_hash':'qa'} for price,currency,transaction in [(100,'EUR','sale'),(90,'USD','sale'),(80,'USD','rent'),(70,'USD','rent'),(None,'USD','rent'),(60,'USD','rent')]]
    ws=load_workbook(io.BytesIO(export_xlsx([p])),data_only=True)['Storico']
    assert ws['D2'].value=='EUR' and ws['D3'].value=='USD'
    assert [ws.cell(row,5).value for row in range(2,8)]==[None,None,None,-10,None,None]


def test_sample_warns_about_single_domain_and_wide_prices(api):
    db,c,s,user,subject,linked,comps=records(api)
    domains=db.all("SELECT id,domain FROM sources WHERE id LIKE 'qa-evidence-%'")
    comp=db.one("SELECT id,price FROM properties WHERE zone='QA · Zona confronto' AND condition='to_renovate' AND id NOT IN (?,?) ORDER BY price DESC LIMIT 1",(subject,linked))
    try:
        db.execute("UPDATE sources SET domain='one.example' WHERE id LIKE 'qa-evidence-%'")
        db.execute('UPDATE properties SET price=900000 WHERE id=?',(comp['id'],))
        groups=c.get('/api/properties/'+subject).json()['market_references']['groups']
        first=groups[0]
        assert first['source_count']==1 and first['asking_delta_pct']==-14.7
        assert 'Una sola fonte nel campione' in first['warnings']
        assert 'Prezzi molto dispersi: confronta i singoli annunci' in first['warnings']
        assert all(g['asking_delta_pct'] is None for g in groups[1:])
    finally:
        for source in domains:db.execute('UPDATE sources SET domain=? WHERE id=?',(source['domain'],source['id']))
        db.execute('UPDATE properties SET price=? WHERE id=?',(comp['price'],comp['id']))


def test_disputed_comparable_is_not_arbitrarily_priced(api):
    db,c,s,user,subject,linked,comps=records(api)
    samples=db.all("SELECT id,availability FROM properties WHERE zone='QA · Zona confronto' AND condition='to_renovate' AND id NOT IN (?,?) ORDER BY price LIMIT 2",(subject,linked))
    a,b=sorted(x['id'] for x in samples)
    db.execute('INSERT INTO duplicate_reviews VALUES(?,?,?,?,?)',(a,b,'same_asset',user,now()))
    try:
        group=c.get('/api/properties/'+subject).json()['market_references']['groups'][0]
        assert group['count']==1 and group['median_sqm'] is None
        assert not any(x['id'] in (a,b) for x in group['items'])
        db.execute("UPDATE properties SET availability='sold' WHERE id=?",(b,))
        group=c.get('/api/properties/'+subject).json()['market_references']['groups'][0]
        assert group['count']==1 and group['median_sqm'] is None
    finally:
        db.execute('DELETE FROM duplicate_reviews WHERE a=? AND b=?',(a,b))
        for sample in samples:db.execute('UPDATE properties SET availability=? WHERE id=?',(sample['availability'],sample['id']))

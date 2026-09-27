import io
from datetime import datetime,timedelta,timezone,date
from uuid import uuid4
import pytest
from openpyxl import load_workbook
from app.db import dump,now
from app.services.decision_support import freshness,DecisionSupport
from app.services.today import queue
from support.decision import decision_dataset


@pytest.mark.parametrize('value',['invalid',None,'2999-01-01T00:00:00+00:00'])
def test_invalid_or_future_acquisition_is_not_recent(value):
    r=freshness({'last_detail_at':value,'config':'{}'})
    assert r['status']=='stale' and r['checked_at'] is None


def test_freshness_uses_source_frequency_and_exact_boundary():
    instant=datetime(2026,9,26,12,tzinfo=timezone.utc)
    row={'kind':'import','config':dump({'detail_refresh_hours':6}),'last_detail_at':(instant-timedelta(hours=6)).isoformat()}
    assert freshness(row,instant)['status']=='stale'
    row['last_detail_at']=(instant-timedelta(hours=5,minutes=59)).isoformat()
    assert freshness(row,instant)['status']=='recent'
    assert freshness(row,instant)['method']=='Importazione'
    row['config']=dump({'detail_refresh_hours':False})
    assert freshness(row,instant)['hours']==24


def dataset(api):
    app,c,s=api;db=app.state.db
    user=c.get('/api/auth/me').json()['user']['id']
    if db.one("SELECT id FROM sources WHERE id='qa-evidence-0'"):
        return db,c,db.one("SELECT id FROM properties WHERE listing_key='Asset da approfondire'")['id'],db.one("SELECT id FROM properties WHERE listing_key='Stesso asset, altra fonte'")['id']
    a,b,_=decision_dataset(db,s,user)
    return db,c,a,b


def test_stale_detail_goes_to_verification_even_with_recent_catalogue(api):
    db,c,a,b=dataset(api)
    old=(datetime.now(timezone.utc)-timedelta(days=3)).isoformat()
    db.execute('UPDATE listing_checks SET last_detail_at=? WHERE property_id IN (?,?)',(old,a,b))
    try:
        rows=queue(db)
        for row in rows['verify']:
            if row['id'] in (a,b):
                assert row['source_name']==c.get('/api/properties/'+row['id']).json()['source_name']
        assert not any(p['id'] in (a,b) for p in rows['call'])
        assert any(p['id'] in (a,b) and 'Disponibilità da ricontrollare' in p['checks'] for p in rows['verify'])
        p=c.get('/api/properties/'+a).json()
        assert p['decision_support']['freshness']['status']=='stale'
        assert any('ancora disponibile' in q for q in p['decision_support']['questions'])
        assert p['availability']=='listed'
    finally:db.execute('UPDATE listing_checks SET last_detail_at=? WHERE property_id IN (?,?)',(now(),a,b))
    assert any(p['id'] in (a,b) for p in queue(db)['call'])


def test_missing_check_is_not_replaced_by_recent_last_seen(api):
    db,c,a,b=dataset(api)
    db.execute('DELETE FROM listing_checks WHERE property_id=?',(a,))
    try:assert DecisionSupport(db,[a]).freshness({'id':a})['status']=='stale'
    finally:db.execute('INSERT INTO listing_checks VALUES(?,?)',(a,now()))


def test_related_contact_visible_without_transferring_mandate(api):
    db,c,a,b=dataset(api);request=str(uuid4())
    body={'request_id':request,'outcome':'documents_requested','contact_name':'QA Broker collegato','note':'Documenti richiesti dopo verifica del mandato.','mandate_status':'confirmed_by_team','next_contact':(date.today()+timedelta(days=2)).isoformat()}
    assert c.post('/api/properties/'+b+'/contacts',json=body).status_code==201
    try:
        detail=c.get('/api/properties/'+a).json();support=detail['decision_support']
        assert support['related_contact']['property_id']==b
        assert detail['decision']['calls']==[]
        assert any('chi ha il mandato' in q for q in support['questions'])
        wb=load_workbook(io.BytesIO(c.get('/api/properties/'+a+'/memo.xlsx').content),data_only=True)
        selection=wb['Selezione'];headers=[cell.value for cell in selection[1]]
        assert selection.cell(2,headers.index('Da chiarire')+1).value=='\n'.join(support['questions'])
        assert wb['Contatti']['I2'].value==b
        assert selection.cell(2,headers.index('Mandato (team)')+1).value is None
        assert not any(p['id'] in (a,b) for p in queue(db)['call'])
    finally:db.execute('DELETE FROM contact_actions WHERE id=?',(request,))


def test_questions_respect_present_declarations_and_own_mandate(api):
    db,c,a,b=dataset(api)
    p=c.get('/api/properties/'+a).json()
    p['decision'].update(cadastral={'quote':'C/1'},published_at='2025-01-01',calls=[{'mandate_status':'confirmed_by_team'}])
    questions=DecisionSupport(db,[a]).summarize(p)['questions']
    assert not any('mandato' in q or 'categoria catastale' in q or 'Da quando' in q for q in questions)
    assert any('prezzo aggiornato' in q for q in questions)

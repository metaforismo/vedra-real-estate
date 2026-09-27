"""Evidence, research and human follow-up contracts; isolated test fixtures only."""
import io
import json
from uuid import uuid4
from datetime import date,timedelta
import pytest
from openpyxl import load_workbook
from app.db import dump,load,now
from app.schemas import Criteria
from app.services.analysis import validate_semantic,prompt_key
from app.services.research_brief import validate_targets,catalog_url
from app.services.decision_facts import dossier,published_date,source_context
from app.services.store import upsert_listing,property_dict
from app.services.today import queue
from app.services.exports import export_xlsx
from app.connectors.parser import extract_listing
from support.catalog import seed,demo_records

PROMPT='Solo cambio d’uso esplicito\nSolo Agenzia Test'
QUOTE='possibile cambio d’uso'


def assessment(checks):
    p={'title':'Locale','description':QUOTE,'source_context':'name: Agenzia Test','custom_prompt':PROMPT,'content_hash':'test'}
    raw={'summary':'Verifica dei criteri richiesti.', 'custom_assessment':{'status':'matched','reason':'La fonte contiene le dichiarazioni richieste.','evidence':[QUOTE],'checks':checks}}
    return validate_semantic(p,raw)['custom_assessments'][prompt_key(PROMPT)]


def checks():
    return [{'criterion':line,'status':'matched','reason':'Dichiarato nella fonte.','evidence':[quote]} for line,quote in zip(PROMPT.splitlines(),(QUOTE,'Agenzia Test'))]


def test_every_requirement_needs_its_own_evidence():
    assert assessment(checks())['status']=='matched'
    partial=assessment(checks()[:1])
    assert partial['status']=='uncertain' and partial['checks'][1]['status']=='uncertain'
    negative=checks();negative[1]['status']='not_matched'
    assert assessment(negative)['status']=='not_matched'
    assert assessment([])['status']=='uncertain'


@pytest.mark.parametrize('field,value',[('criterion','Non richiesto'),('evidence',['Mandato garantito']),('evidence',[])])
def test_per_criterion_rejects_fabricated_or_unproven_checks(field,value):
    items=checks();items[0][field]=value
    with pytest.raises(ValueError):assessment(items)


def test_multiline_normalization_and_source_target_validation():
    c=Criteria(custom_prompt='  Primo   criterio\n\n Secondo criterio ',research_instructions=' Visita prima il broker\n Poi il catalogo')
    assert c.custom_prompt=='Primo criterio\nSecondo criterio'
    source={'id':'web','kind':'html','domain':'catalog.example','config':dump({'search_url':'https://catalog.example/{city}'})}
    validate_targets(Criteria(source_urls={'web':'https://catalog.example/broker/one'}),[source])
    assert catalog_url({'city':'Reggio Calabria','criteria':{}},source).endswith('reggio-calabria')
    for url in ('https://elsewhere.example/','https://user@catalog.example/','https://catalog.example:8080/','javascript:alert(1)','https://catalog.example\\@elsewhere.example/'):
        with pytest.raises(ValueError):validate_targets(Criteria(source_urls={'web':url}),[source])
    with pytest.raises(ValueError):validate_targets(Criteria(source_urls={'unknown':'https://catalog.example/'}),[source])


def html(contact=True):
    node={'@type':'Apartment','name':'Locale test','description':'Categoria catastale C/1. Cambio di destinazione d’uso non consentito. Mandato in esclusiva dichiarato.', 'datePublished':'2025-01-02', 'offers':{'price':550000,'priceCurrency':'EUR'},'floorSize':{'value':100}}
    if contact:node['offers']['seller']={'@type':'Organization','name':'Agenzia Test','telephone':'+39 02 1234567','email':'info@catalog.example'}
    return '<script type="application/ld+json">'+json.dumps([node,{'@type':'Organization','name':'Footer estraneo','telephone':'12345678'}])+'</script>'


def test_contact_and_declarations_preserve_provenance_and_negation():
    p=extract_listing(html(),'https://catalog.example/listing/one')
    facts=p.evidence['decision_facts']
    assert facts['contact']['telephone']=='+39021234567'
    assert facts['contact']['role']=='Inserzionista dichiarato'
    assert 'C/1' in facts['cadastral']['quote'] and 'non consentito' in facts['change_of_use']['quote']
    assert facts['mandate']['status']=='dichiarato nella fonte'
    assert facts['published_at'].startswith('2025-01-02')
    assert 'contact' not in extract_listing(html(False),'https://catalog.example/listing/one').evidence['decision_facts']
    assert published_date('2999-01-01') is None and published_date('invalid') is None
    assert dossier(p.model_dump(),[{'price':v,'currency':'EUR','transaction_type':'sale'} for v in (100,None,90,80)])['price_reductions']==1


def test_broker_changes_invalidate_analysis_input(db,settings):
    seed(db,settings)
    p=extract_listing(html(),'https://catalog.example/listing/one')
    ident,_,_=upsert_listing(db,settings,'demo-milano',p)
    before=property_dict(db.one('SELECT * FROM properties WHERE id=?',(ident,)))
    p=extract_listing(html().replace('Agenzia Test','Agenzia Nuova'),'https://catalog.example/listing/one')
    upsert_listing(db,settings,'demo-milano',p)
    after=property_dict(db.one('SELECT * FROM properties WHERE id=?',(ident,)))
    assert before['content_hash']!=after['content_hash']
    assert 'Agenzia Nuova' in source_context(after)


def test_contact_ledger_validation_idempotency_queue_and_excel(api):
    app,c,settings=api;db=app.state.db
    p=extract_listing(html(),'https://catalog.example/listing/contact-workflow')
    ident,_,_=upsert_listing(db,settings,'demo-milano',p)
    db.execute("UPDATE properties SET availability='listed',review_status='shortlisted' WHERE id=?",(ident,))
    assert any(x['id']==ident for x in queue(db)['call'])
    body={'request_id':str(uuid4()),'contact_name':'Agenzia Test','outcome':'no_answer','mandate_status':'not_checked','note':'=Test di escaping'}
    route=f'/api/properties/{ident}/contacts'
    assert c.post(route,json=body|{'mandate_status':'confirmed_by_team','note':''}).status_code==422
    assert c.post(route,json=body|{'outcome':'not_relevant','next_contact':'2026-01-01'}).status_code==422
    for _ in range(2):assert c.post(route,json=body).status_code==201
    assert c.post(route,json=body|{'note':'changed'}).status_code==409
    detail=c.get('/api/properties/'+ident).json()
    assert len(detail['decision']['calls'])==1
    wb=load_workbook(io.BytesIO(export_xlsx([detail],db)),data_only=True)
    sheet=wb['Selezione'];headers=[cell.value for cell in sheet[1]]
    assert sheet.cell(2,headers.index('Inserzionista')+1).value=='Agenzia Test'
    assert wb['Contatti']['F2'].value=="'=Test di escaping"
    assert wb['Contatti']['F2'].data_type=='s'
    tomorrow=(date.today()+timedelta(days=1)).isoformat()
    follow=body|{'request_id':str(uuid4()),'outcome':'documents_requested','next_contact':tomorrow}
    assert c.post(route,json=follow).status_code==201
    assert all(x['id']!=ident for x in queue(db)['call'])
    db.execute('UPDATE contact_actions SET next_contact=? WHERE id=?',(date.today().isoformat(),follow['request_id']))
    assert any(x['id']==ident for x in queue(db)['call'])
    db.execute("UPDATE properties SET availability='sold' WHERE id=?",(ident,))
    assert all(x['id']!=ident for x in queue(db)['call'])
    assert c.post('/api/properties/no-such-property/contacts',json=body|{'request_id':str(uuid4())}).status_code==404
    user=c.get('/api/auth/me').json()['user']['id']
    db.execute("UPDATE users SET role='viewer' WHERE id=?",(user,))
    try:assert c.post(route,json=body|{'request_id':str(uuid4())}).status_code==403
    finally:db.execute("UPDATE users SET role='admin' WHERE id=?",(user,))


def test_agent_research_requires_online_hermes_and_selected_source(api):
    app,c,_=api;db=app.state.db
    db.execute('INSERT INTO sources(id,name,kind,domain,config,permission_at,created_at) VALUES(?,?,?,?,?,?,?)',('research-web','Research test','html','catalog.example',dump({'search_url':'https://catalog.example/search'}),now(),now()))
    body={'name':'Research instructions','city':'Milano','runtime':'hermes','source_ids':['research-web'],'criteria':{'research_instructions':'Visita prima la pagina del broker','source_urls':{'research-web':'https://catalog.example/broker'}}}
    assert c.post('/api/agents',json=body).status_code==422
    body['criteria']['online_discovery']=True
    response=c.post('/api/agents',json=body);assert response.status_code==201,response.text
    agent=next(a for a in c.get('/api/agents').json() if a['id']==response.json()['id'])
    assert agent['criteria']['research_instructions']==body['criteria']['research_instructions']
    assert agent['criteria']['source_urls']==body['criteria']['source_urls']


def test_price_cut_does_not_compare_currency_or_transaction_changes():
    p={'url':'https://catalog.example/1'}
    rows=[{'price':100,'currency':'EUR','transaction_type':'sale'}, {'price':90,'currency':'USD','transaction_type':'sale'}, {'price':80,'currency':'USD','transaction_type':'rent'}, {'price':70,'currency':'USD','transaction_type':'rent'}]
    assert dossier(p,rows)['price_reductions']==1
    assert dossier(p,[{'price':100},{'price':90}])['price_reductions']==0


def test_unversioned_assets_revalidate(api):
    _,c,_=api
    response=c.get('/assets/product-actions.js')
    assert response.status_code==200 and response.headers['Cache-Control']=='no-cache'

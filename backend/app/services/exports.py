from __future__ import annotations

import csv
import io
from datetime import datetime,timezone

from ..security import csv_safe

HEADERS=['ID','Titolo','Comune','Micro-zona','Tipologia','Prezzo','Superficie mq','Prezzo/mq',
         'Benchmark min','Benchmark max','Sconto %','Score preliminare','Completezza %','Motore','Dataset','Fonte','URL','Rilevato il','Valuta']


def export_csv(rows):
    output=io.StringIO(newline='')
    writer=csv.writer(output,delimiter=';')
    writer.writerow(HEADERS)
    for p in rows:
        b=p.get('benchmark') or {}
        writer.writerow([csv_safe(v) for v in [p['id'],p['title'],p['city'],p['zone'],p['property_type'],p['price'],p['surface'],
            p['price_sqm'],b.get('min_sqm'),b.get('max_sqm'),p['discount'],p['score'],p['completeness'],p['analysis'].get('engine'),
            'SINTETICO / DEMO' if p['is_demo'] else 'REALE',p.get('source_name',''),p['url'],p['last_seen'],p['currency']]])
    return ('\ufeff'+output.getvalue()).encode('utf-8')


def export_xlsx(rows):
    """Portable application exporter; no proprietary runtime required on the user's VPS."""
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.comments import Comment
    from openpyxl.worksheet.table import Table,TableStyleInfo
    wb=Workbook()
    ws=wb.active;ws.title='Opportunità'
    ws.append(['VEDRA | Screening preliminare, non perizia. I dati DEMO sono sintetici.'])
    ws.merge_cells('A1:S1');ws.row_dimensions[1].height=32
    ws['A1'].font=Font(size=14,bold=True,color='173C35')
    ws.append(HEADERS)
    for i,p in enumerate(rows,3):
        b=p.get('benchmark') or {}
        cells=[p['id'],p['title'],p['city'],p['zone'],p['property_type'],p['price'],p['surface'],
              f'=IF(OR(F{i}="",G{i}="",G{i}=0),"",F{i}/G{i})',b.get('min_sqm'),b.get('max_sqm'),
              f'=IF(OR(H{i}="",I{i}="",J{i}=""),"",1-H{i}/AVERAGE(I{i}:J{i}))',p['score'],p['completeness'],
              p['analysis'].get('engine'), 'SINTETICO / DEMO' if p['is_demo'] else 'REALE',p.get('source_name',''),p['url'],p['last_seen'],p['currency']]
        for j,val in enumerate(cells,1):
            if isinstance(val,str) and j not in (8,11):val=csv_safe(val)
            ws.cell(i,j,val)
        for j in (6,8,9,10):ws.cell(i,j).number_format=f'#,##0.00 "{p["currency"] if p["currency"]!="XXX" else "valuta n.d."}"'
        ws.cell(i,7).number_format='#,##0.0'
        ws.cell(i,11).number_format='0.0%'
        ws.cell(i,12).number_format='0" / 100"'
        ws.cell(i,13).number_format='0"%"'
        ws.cell(i,9).comment=Comment(f"Fonte: {b.get('source_label','Nessun benchmark')}\n{b.get('source_url','')}\nPeriodo: {b.get('period','')}\nConfronto sul punto medio, non su un prezzo transato.",'Vedra')
        ws.row_dimensions[i].height=30
    for c in ws[2]:
        c.fill=PatternFill('solid',fgColor='173C35');c.font=Font(color='FFFFFF',bold=True);c.alignment=Alignment(wrap_text=True,vertical='center')
    ws.row_dimensions[2].height=30
    from openpyxl.utils import get_column_letter
    for j in range(1,20):ws.column_dimensions[get_column_letter(j)].width=18
    for col,width in [('A',38),('B',40),('D',23),('P',30),('Q',44),('R',27)]:ws.column_dimensions[col].width=width
    for row in ws.iter_rows(min_row=3):
        for cell in row:cell.alignment=Alignment(vertical='center',wrap_text=True)
    ws.freeze_panes='F3'
    if rows:
        table=Table(displayName='Opportunita',ref=f'A2:S{len(rows)+2}')
        table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True)
        ws.add_table(table)
    note=wb.create_sheet('Metodo e limiti')
    notes=[['VEDRA | Metodo'],['Scope','Origination e screening preliminare. Nessun rendimento o cambio d’uso certificato.'],
           ['Fonte','I benchmark vanno importati dal cliente. Non sono scaricati automaticamente da OMI.'],
           ['Confronto','Richiede comune, micro-zona, tipo, stato, valuta, contratto e base superficie compatibili.'],
           ['Completezza','Campi presenti / 10 campi attesi. Non è accuratezza, copertura del mercato o probabilità.'],
           ['Score','clamp(35 + sconto % × 1,4; 0; 70) + min(30; 15 × strategie con evidenza).'],
           ['Assenza benchmark','Score e delta restano vuoti. Non sostituiti con una media di città.'],
           ['Formule','Prezzo/mq e delta si ricalcolano in Excel. Nessun prezzo di uscita o margine inventato.'],
           ['Esportazione',datetime.now(timezone.utc).isoformat(timespec='seconds')]]
    for r in notes:note.append(r)
    note.column_dimensions['A'].width=25;note.column_dimensions['B'].width=95
    for row in note:
        for cell in row:cell.alignment=Alignment(wrap_text=True,vertical='top')
        note.row_dimensions[row[0].row].height=35
    # Excel recalculates on opening. Also populate the two deterministic caches so
    # read-only viewers do not display empty values before their first recalculation.
    wb.calculation.fullCalcOnLoad=True
    wb.calculation.calcMode='auto'
    out=io.BytesIO();wb.save(out)
    import zipfile
    from xml.etree import ElementTree as ET
    namespace={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    patched=io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(out.getvalue())) as source,zipfile.ZipFile(patched,'w',zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            content=source.read(item.filename)
            if item.filename=='xl/worksheets/sheet1.xml':
                tree=ET.fromstring(content)
                cells={cell.get('r'):cell for cell in tree.findall('.//s:c',namespace)}
                for i,p in enumerate(rows,3):
                    amount=p['price']/p['surface'] if p['price'] and p['surface'] else None
                    b=p.get('benchmark')
                    delta=1-amount/((b['min_sqm']+b['max_sqm'])/2) if b and amount is not None else None
                    for address,value in ((f'H{i}',amount),(f'K{i}',delta)):
                        cell=cells[address]
                        cached=cell.find('s:v',namespace)
                        if cached is None:cached=ET.SubElement(cell,'{'+namespace['s']+'}v')
                        cached.text=str(value) if value is not None else ''
                        cell.set('t','n' if value is not None else 'str')
                content=ET.tostring(tree,encoding='utf-8',xml_declaration=True)
            target.writestr(item,content)
    return patched.getvalue()


def export_docx(p):
    from docx import Document
    from docx.shared import Inches,Pt,RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc=Document()
    sec=doc.sections[0]
    sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.8);sec.right_margin=Inches(.8)
    normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10)
    normal.paragraph_format.space_after=Pt(6)
    doc.core_properties.title=f"Vedra · {p['title']}";doc.core_properties.author='Vedra'
    sec.header.paragraphs[0].text='VEDRA  /  REAL ESTATE INTELLIGENCE'
    doc.add_heading('Scheda di screening',0)
    doc.add_paragraph('DATI SINTETICI / DEMO' if p['is_demo'] else 'DOCUMENTO PRELIMINARE · VERIFICA UMANA RICHIESTA','Subtitle')
    doc.add_heading(p['title'],1)
    doc.add_paragraph(f"{p['city']} · {p['zone'] or 'Micro-zona non disponibile'}\nID: {p['id']}")
    def euro(v):return (f'{v:,.0f}'.replace(',','.')+' '+('€' if p['currency']=='EUR' else p['currency'] if p['currency']!='XXX' else '(valuta n.d.)')) if v is not None else 'Non disponibile'
    table=doc.add_table(rows=0,cols=2);table.style='Light Shading Accent 1'
    for k,v in [('Prezzo richiesto',euro(p['price'])),('Superficie',f"{p['surface']} mq" if p['surface'] else 'Non disponibile'),
                ('Prezzo / mq',euro(p['price_sqm'])),('Score di screening',str(p['score'])+'/100' if p['score'] is not None else 'Non disponibile'),
                ('Completezza campi',f"{p['completeness']}% (non misura l’accuratezza)"),('Tipologia / stato',f"{p['property_type']} / {p['condition']}")]:
        row=table.add_row();row.cells[0].text=k;row.cells[1].text=v
    doc.add_heading('Strategie ed evidenze',2)
    strategies=p['analysis'].get('strategies',[])
    if not strategies:doc.add_paragraph('Nessuna strategia supportata dal testo acquisito.')
    for s in strategies:
        doc.add_paragraph(s['strategy'].replace('_',' ').title(),'Heading 3')
        doc.add_paragraph('Evidenza dal testo: “'+s['evidence']+'”')
    b=p.get('benchmark')
    doc.add_heading('Confronto di prezzo',2)
    if b:
        doc.add_paragraph(f"Range: {euro(b['min_sqm'])} – {euro(b['max_sqm'])} / mq. Periodo {b['period']}.\nScostamento del prezzo dal punto medio: {-p['discount']}% (negativo = sotto il riferimento). Non è una stima del prezzo transato.")
        doc.add_paragraph(f"Fonte benchmark: {b['source_label']}\n{b['source_url']}")
    else:doc.add_paragraph(p['analysis'].get('benchmark_note','Nessun benchmark compatibile.'))
    doc.add_heading('Fonti e limiti',2)
    doc.add_paragraph(f"Annuncio: {p['url']}\nAcquisito: {p['last_seen']}\nMotore classificazione: {p['analysis'].get('engine','non disponibile')}")
    doc.add_paragraph('Campi assenti: '+(', '.join(p['missing_fields']) or 'nessuno tra quelli misurati'))
    for text in p['analysis'].get('caveats',[]):doc.add_paragraph(text)
    doc.add_paragraph('Questo documento non costituisce una perizia o una raccomandazione d’investimento. Destinazione d’uso, urbanistica, stato locativo, costi e diritti sui dati devono essere verificati prima di ogni decisione.')
    footer=sec.footer.paragraphs[0];footer.text='VEDRA · Preview 0.1  |  '
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
    out=io.BytesIO();doc.save(out);return out.getvalue()

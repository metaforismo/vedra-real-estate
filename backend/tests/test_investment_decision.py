"""Direct leads and decision economics: explicit evidence and reproducible arithmetic."""
import io
import pytest
from openpyxl import load_workbook
from app.product_schemas import ScenarioInputs
from app.schemas import Criteria
from app.services.scenarios import calculate
from app.services.analysis import screen
from app.services.decision_facts import contact_route
from app.services.research_brief import brief


@pytest.mark.parametrize('text,kind',[
    ('Vendita diretta dal proprietario. Appartamento libero.','owner_declared'),
    ('Vendita diretta da privato.','owner_declared'),
    ('Vendo da privato.','owner_declared'),
    ('Agenzia con mandato in esclusiva.','mandate_declared'),
    ('Incarico in esclusiva per questo immobile.','mandate_declared'),
    ('Non abbiamo un mandato esclusivo.','unknown'),
    ('Mandato esclusivo scaduto.','unknown'),
    ('Cerchiamo immobili con mandato esclusivo.','unknown'),
    ('Vendita diretta da un intermediario.','unknown'),
    ('No agenzie, ottimo investimento.','unknown'),
    ('Vendita diretta dal proprietario. Mandato esclusivo.','unknown'),
])
def test_direct_route_needs_specific_noncontradictory_declaration(text,kind):
    assert contact_route({'title':'Appartamento','description':text,'url':'https://qa.example/1'})['kind']==kind


def candidate():
    return {'title':'Asset','description':'Vendita diretta dal proprietario.','url':'https://qa.example/1',
            'city':'Milano','availability':'listed','transaction_type':'sale','currency':'EUR','price':200000,'surface':100,
            'discount':12,'analysis':{},'evidence':{'decision_facts':{'contact':{'telephone':'+39020000000'}}}}


def test_opportunity_filter_does_not_promote_missing_or_market_aligned_prices():
    agent={'city':'Milano','criteria':{'opportunity_only':True,'contact_policy':'require_direct','min_discount':10}}
    p=candidate();assert screen(p,agent)[0]
    for discount in (None,-10,0,5):assert not screen(p|{'discount':discount},agent)[0]
    assert not screen(p|{'description':'No agenzie'},agent)[0]
    assert not screen(p|{'evidence':{}},agent)[0]
    assert screen(p|{'description':'Non specificato'},{'city':'Milano','criteria':{'contact_policy':'prefer_direct'}})[0]


def test_research_brief_transmits_structured_constraints():
    criteria=Criteria(opportunity_only=True,contact_policy='require_direct',min_discount=12).model_dump()
    b=brief({'city':'Milano','criteria':criteria},[])
    assert b['opportunity_only'] and b['contact_policy']=='require_direct' and b['min_discount']==12
    assert 'pagina' in b['contact_task']
    assert b['version']!=brief({'city':'Milano','criteria':criteria|{'contact_policy':'any'}},[])['version']


def assumptions(**extra):
    return ScenarioInputs(purchase=200000,sale=320000,works=40000,acquisition_costs=20000,
                          contingency_pct=10,selling_pct=3,holding_monthly=1000,months=12,target_roi_pct=20,**extra)


def test_combined_stress_ceiling_and_target_are_consistent():
    r=calculate(assumptions())
    assert r['invested']==276000 and r['profit']==34400
    assert r['max_purchase']==pytest.approx(182666.66)
    assert not r['target_met']
    assert r['stress']['invested']==290800 and r['stress']['profit']==-11440
    assert r['stress']['max_purchase']==142000
    assert r['stress']['roi_pct']<r['roi_pct']
    changed=assumptions().model_copy(update={'purchase':r['max_purchase']})
    assert calculate(changed)['roi_pct']==20 and calculate(changed)['target_met']


def test_zero_stress_matches_base_and_negative_ceiling_is_not_clamped():
    r=calculate(assumptions(stress_sale_pct=0,stress_works_pct=0,stress_delay_months=0))
    assert all(r[k]==r['stress'][k] for k in ('invested','profit','roi_pct','max_purchase','breakeven_sale','target_met'))
    assert calculate(ScenarioInputs(purchase=100,sale=10,works=1000))['max_purchase']<0


@pytest.mark.parametrize('key,value',[('target_roi_pct',-1),('stress_sale_pct',100),('stress_works_pct',301),('stress_delay_months',1.5),('sale',float('inf'))])
def test_bad_scenarios_rejected(key,value):
    from pydantic import ValidationError
    with pytest.raises(ValidationError):ScenarioInputs.model_validate({'purchase':100,'sale':200,key:value})


def test_saved_scenario_matches_platform_and_excel(api):
    app,c,s=api
    pid=app.state.db.one("SELECT id FROM properties WHERE is_demo=0 AND currency='EUR'")['id']
    inputs=assumptions().model_dump();body={'name':'QA decision','inputs':inputs}
    saved=c.post(f'/api/properties/{pid}/scenarios',json=body)
    assert saved.status_code==201
    expected=saved.json()['result']
    try:
        read=c.get(f'/api/properties/{pid}/scenarios').json()
        row=next(x for x in read if x['id']==saved.json()['id'])
        assert row['result']==expected and row['inputs']==inputs
        response=c.get(f'/api/properties/{pid}/memo.xlsx')
        wb=load_workbook(io.BytesIO(response.content),data_only=True)
        ws=wb['Scenari'];values=next(row for row in ws.iter_rows(min_row=2,values_only=True) if row[1]=='QA decision')
        assert values[16]==expected['invested'] and values[17]==expected['profit']
        assert values[20]==expected['max_purchase'] and values[22]==expected['stress']['profit']
        assert values[24]==expected['stress']['max_purchase'] and values[26]==expected['version']
    finally:c.delete('/api/scenarios/'+saved.json()['id'])


def test_same_asset_prefers_direct_lead_over_higher_priority_middleman(api):
    from support.decision import decision_dataset
    from app.services.today import queue
    app,c,s=api;db=app.state.db
    user=c.get('/api/auth/me').json()['user']['id']
    subject,linked,_=decision_dataset(db,s,user)
    db.execute("UPDATE properties SET priority_score=100 WHERE id=?",(subject,))
    db.execute("UPDATE properties SET description='Vendita diretta dal proprietario.',priority_score=0 WHERE id=?",(linked,))
    rows=[x for x in queue(db)['call'] if x['id'] in (subject,linked)]
    assert len(rows)==1 and rows[0]['id']==linked
    assert rows[0]['contact_route']['kind']=='owner_declared'

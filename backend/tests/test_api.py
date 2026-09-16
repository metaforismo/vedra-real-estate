import io
from fastapi.testclient import TestClient
from app.db import load


def test_auth_required_and_headers(api):
    app,client,_=api
    anon=TestClient(app)
    assert anon.get('/api/workspace').status_code==401
    response=client.get('/')
    assert response.status_code==200 and "script-src 'self'" in response.headers['content-security-policy']
    assert response.headers['x-content-type-options']=='nosniff'
    assert client.get('/api/workspace').headers['cache-control']=='no-store'
    assert client.get('/api/docs').status_code==200
    assert client.get('/api/openapi.json').json()['info']['version']=='0.3.0'


def test_fixture_counts_and_separation(api):
    _,c,_=api
    data=c.get('/api/workspace?dataset=real').json()
    assert data['stats']['properties']==36 and len(data['agents'])==3
    assert not any(p['is_demo'] for p in data['properties'])
    assert all('sintetico' in p['benchmark']['source_label'] for p in data['properties'] if p['benchmark'])
    assert data['stats']['benchmarked']==36
    assert len(c.get('/api/workspace?dataset=real').json()['properties'])==36
    assert c.get('/api/workspace?dataset=demo').status_code==422
    assert c.get('/api/workspace?dataset=evil').status_code==422


def test_csrf_and_origin(api):
    _,c,_=api
    csrf=c.headers.pop('x-csrf-token')
    assert c.post('/api/agents/agent-milano/toggle').status_code==403
    c.headers['X-CSRF-Token']=csrf
    assert c.post('/api/agents/agent-milano/toggle',headers={'Origin':'https://evil.example'}).status_code==403


def test_crud_agent_and_queue_idempotence(api):
    app,c,_=api
    payload={'name':'Ricerca test','city':'Milano','source_ids':['demo-milano'],'criteria':{'max_price':700000,'min_surface':200},'interval_minutes':0}
    response=c.post('/api/agents',json=payload);assert response.status_code==201,response.text
    aid=response.json()['id']
    first=c.post(f'/api/agents/{aid}/run').json();second=c.post(f'/api/agents/{aid}/run').json()
    assert first['id']==second['id'] and first['status']=='queued'
    payload['criteria']['max_price']=300000
    assert c.put(f'/api/agents/{aid}',json=payload).status_code==200
    run=c.get('/api/runs/'+first['id']).json()
    assert run['config_snapshot']['criteria']['max_price']==700000
    c.post('/api/runs/'+first['id']+'/cancel')
    assert c.get('/api/runs/'+first['id']).json()['status']=='cancelled'
    assert c.post(f'/api/agents/{aid}/toggle').json()['active'] is False


def test_legacy_demo_sources_cannot_be_selected(api):
    app,c,_=api
    from app.db import now
    app.state.db.execute("INSERT INTO sources(id,name,kind,created_at) VALUES('legacy-demo','Legacy','demo',?)",(now(),))
    response=c.post('/api/agents',json={'name':'Misto','city':'Milano','source_ids':['legacy-demo']})
    assert response.status_code==422
    assert c.get('/api/workspace?dataset=demo').status_code==422
    assert c.get('/api/workspace?dataset=all').status_code==422
    assert 'legacy-demo' not in [x['id'] for x in c.get('/api/sources').json()]


def test_notes_review_snapshot_and_events(api):
    _,c,_=api
    p=c.get('/api/workspace?dataset=real').json()['properties'][0]
    assert c.post(f"/api/properties/{p['id']}/notes",json={'body':'Verificare la superficie commerciale.'}).status_code==201
    assert c.patch(f"/api/properties/{p['id']}",json={'review_status':'reviewing','starred':True}).status_code==200
    detail=c.get(f"/api/properties/{p['id']}").json()
    assert detail['notes'][0]['body'].startswith('Verificare') and detail['review_status']=='reviewing'
    snap=c.get(f"/api/properties/{p['id']}/snapshot")
    assert snap.status_code==200 and 'attachment' in snap.headers['content-disposition']
    assert 'DATI SINTETICI' in snap.text
    rid=c.get('/api/workspace?dataset=real').json()['runs'][-1]['id']
    events=c.get(f'/api/runs/{rid}/events')
    assert events.status_code==200 and 'event: done' in events.text


def test_viewer_cannot_write_and_analyst_cannot_admin(api):
    app,admin,settings=api
    for role in ('viewer','analyst'):
        response=admin.post('/api/users',json={'name':'Test '+role,'email':role+'@example.test','password':'test-account-password-92842','role':role})
        assert response.status_code==201
        c=TestClient(app)
        login=c.post('/api/auth/login',json={'email':role+'@example.test','password':'test-account-password-92842'})
        assert login.status_code==200
        c.headers['X-CSRF-Token']=login.json()['csrf']
        assert c.get('/api/workspace').status_code==200
        assert c.get('/api/users').status_code==403
        assert c.get('/api/runtime').status_code==403
        aid='agent-como'
        if role=='viewer':assert c.post(f'/api/agents/{aid}/toggle').status_code==403
        else:
            assert c.post(f'/api/agents/{aid}/toggle').status_code==200
            c.post(f'/api/agents/{aid}/toggle')
        assert c.post('/api/auth/logout').status_code==200
        assert c.get('/api/auth/me').status_code==401


def test_large_body_rejected_before_parsing(api):
    _,c,_=api
    assert c.post('/api/imports',content=b'x'*5_000_001).status_code==413


def test_exports_actual_office_documents(api):
    from openpyxl import load_workbook
    from docx import Document
    _,c,_=api
    p=c.get('/api/workspace?dataset=real').json()['properties'][0]
    response=c.post('/api/export',json={'format':'xlsx','dataset':'real','ids':[p['id']]})
    assert response.status_code==200,response.text
    wb=load_workbook(io.BytesIO(response.content),data_only=False)
    assert wb['Opportunità']['H3'].value.startswith('=IF(')
    assert wb['Opportunità']['K3'].value.startswith('=IF(')
    assert wb['Opportunità']['O3'].value=='REALE'
    assert wb['Opportunità']['S3'].value=='EUR'
    values=load_workbook(io.BytesIO(response.content),data_only=True)
    assert abs(values['Opportunità']['H3'].value-p['price']/p['surface'])<0.0001
    assert abs(values['Opportunità']['K3'].value-p['discount']/100)<0.0001
    response=c.get(f"/api/properties/{p['id']}/memo.docx")
    assert response.status_code==200
    doc=Document(io.BytesIO(response.content))
    text=' '.join(x.text for x in doc.paragraphs)
    assert 'DATI SINTETICI' not in text and p['url'] in text and 'perizia' in text
    response=c.post('/api/export',json={'format':'csv','dataset':'real','ids':[p['id']]})
    assert response.content.startswith(b'\xef\xbb\xbf') and 'SINTETICO / DEMO' not in response.text


def test_no_credentials_in_workspace(api):
    _,c,s=api
    text=c.get('/api/workspace').text
    assert s.hermes_key not in text and s.bridge_token not in text and s.admin_password not in text


def test_runtime_discovery_without_network_failure_leak(api):
    app,c,s=api
    key=s.hermes_key;s.hermes_key=''
    try:
        data=c.get('/api/runtime').json()
        assert not data['hermes_configured']
    finally:s.hermes_key=key

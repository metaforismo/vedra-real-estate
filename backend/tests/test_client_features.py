"""Client workflow contracts; all records and model responses are test fixtures."""
import io
import pytest
from openpyxl import load_workbook
from app.db import dump, load
from app.schemas import Criteria
from app.services.analysis import custom_assessment, prompt_key, screen, validate_semantic
from app.services.store import property_dict, refresh_analysis, upsert_listing
from app.services.engine import Engine
from app.services.market_references import MarketReferences
from app.services.exports import export_xlsx
from support.catalog import seed, demo_records

PROMPT='Solo immobili con cambio d’uso esplicito'
QUOTE='possibile cambio d’uso'


def listing(db,settings):
    seed(db,settings)
    return property_dict(db.one("SELECT * FROM properties WHERE source_id='demo-milano' AND title LIKE 'Spazi direzionali%'"))


def result(p,prompt=PROMPT,status='matched',evidence=None):
    return validate_semantic(p|{'custom_prompt':prompt},{'summary':'Test di interpretazione qualitativa.','strategies':[], 'caveats':[], 'custom_assessment':{'status':status,'reason':'Il testo dichiara la possibilità, soggetta a verifica tecnica.','evidence':[QUOTE] if evidence is None else evidence}})


def test_custom_criteria_cannot_override_filters_and_unknown_is_not_fit(db,settings):
    p=listing(db,settings)
    agent={'city':'Milano','criteria':{'custom_prompt':PROMPT,'max_price':2000000}}
    assert not screen(p,agent)[0]
    p['analysis']=result(p)
    assert screen(p,agent)[0]
    assert not screen(p|{'price':3000000},agent)[0]
    assert not screen(p|{'discount':None},agent|{'criteria':agent['criteria']|{'min_discount':20}})[0]
    for status in ('not_matched','uncertain'):
        p['analysis']=result(p,status=status)
        assert not screen(p,agent)[0]


def test_custom_scoped_to_prompt_and_listing_version(db,settings):
    p=listing(db,settings);p['analysis']=result(p)
    assert custom_assessment(p,PROMPT)
    assert not custom_assessment(p,PROMPT+' e piano terra')
    assert not custom_assessment(p|{'content_hash':'changed'},PROMPT)
    assert Criteria(custom_prompt='  solo   piano terra ').custom_prompt=='solo piano terra'
    with pytest.raises(ValueError):Criteria(custom_prompt='x'*6001)


@pytest.mark.parametrize('evidence',[[],['Autorizzazione definitiva garantita']])
def test_custom_definitive_requires_real_quotes(db,settings,evidence):
    with pytest.raises(ValueError):result(listing(db,settings),evidence=evidence)


def test_custom_uncertain_can_have_no_quote_but_missing_assessment_rejected(db,settings):
    p=listing(db,settings)
    assert result(p,status='uncertain',evidence=[])['custom_assessments']
    with pytest.raises(ValueError):validate_semantic(p|{'custom_prompt':PROMPT},{'summary':'Risposta senza verifica'})


def test_multiple_agents_checks_survive_refresh_and_change_discards(db,settings):
    p=listing(db,settings)
    refresh_analysis(db,p['id'],result(p))
    other='Solo piano terra'
    updated=refresh_analysis(db,p['id'],result(p,prompt=other,status='uncertain',evidence=[]))
    assert custom_assessment(updated,PROMPT)['status']=='matched'
    assert custom_assessment(updated,other)['status']=='uncertain'
    source,raw=demo_records('Milano')[0]
    upsert_listing(db,settings,'demo-milano',source.model_copy(update={'description':source.description+' Nuovo testo.'}),raw=raw)
    updated=property_dict(db.one('SELECT * FROM properties WHERE id=?',(p['id'],)))
    assert not custom_assessment(updated,PROMPT)


async def test_new_prompt_queues_semantic_task_on_unchanged_listing(db,settings):
    p=listing(db,settings);engine=Engine(db,settings)
    refresh_analysis(db,p['id'],result(p))
    agent=db.one("SELECT * FROM agents WHERE id='agent-milano'")
    criteria=load(agent['criteria'])|{'custom_prompt':PROMPT}
    db.execute("UPDATE agents SET criteria=?,runtime='hermes' WHERE id='agent-milano'",(dump(criteria),))
    run=engine.enqueue('agent-milano');await engine.collect(run['id'])
    assert not db.one('SELECT * FROM semantic_tasks WHERE run_id=? AND property_id=?',(run['id'],p['id']))
    db.execute("UPDATE runs SET status='completed' WHERE id=?",(run['id'],))
    criteria['custom_prompt']=PROMPT+' e piano terra'
    db.execute("UPDATE agents SET criteria=? WHERE id='agent-milano'",(dump(criteria),))
    run=engine.enqueue('agent-milano');await engine.collect(run['id'])
    task=db.one('SELECT * FROM semantic_tasks WHERE run_id=? AND property_id=?',(run['id'],p['id']))
    payload=load(task['payload'])
    assert payload['custom_prompt']==criteria['custom_prompt'] and payload['content_hash']==p['content_hash']


def test_agent_prompt_roundtrip_and_engine_validation(api):
    app,c,_=api
    body={'name':'Criteri qualitativi test','city':'Milano','runtime':'local','source_ids':['demo-milano'], 'criteria':{'custom_prompt':PROMPT}}
    assert c.post('/api/agents',json=body).status_code==422
    response=c.post('/api/agents',json=body|{'runtime':'hermes'})
    assert response.status_code==201,response.text
    ident=response.json()['id']
    a=next(a for a in c.get('/api/agents').json() if a['id']==ident)
    assert a['criteria']['custom_prompt']==PROMPT
    assert c.put('/api/agents/'+ident,json=body|{'runtime':'hermes','criteria':{'custom_prompt':'Solo piano terra'}}).status_code==200
    a=next(a for a in c.get('/api/agents').json() if a['id']==ident)
    assert a['criteria']['custom_prompt']=='Solo piano terra'


def populate_comparables(db,settings):
    p=listing(db,settings)
    source,raw=demo_records('Milano')[0]
    ids=[]
    for condition in ('to_renovate','renovated','new','good'):
        for i in range(3):
            key=f'{condition}-{i}'
            comparable=source.model_copy(update={'listing_key':key,'url':f'https://catalog.example/{key}', 'condition':condition,'price':510*(2000+i*500),'availability':'listed'})
            ident,_,_=upsert_listing(db,settings,'demo-milano',comparable,raw=raw)
            ids.append(ident)
    return p,ids


def test_references_separate_categories_and_dont_invent(db,settings):
    p,ids=populate_comparables(db,settings)
    refs=MarketReferences(db).for_property(p)
    assert [g['count'] for g in refs['groups']]==[3,3,3]
    assert [g['median_sqm'] for g in refs['groups']]==[2500,2500,2500]
    assert all(item['url'].startswith('https://catalog.example/') for g in refs['groups'] for item in g['items'])
    db.execute("UPDATE properties SET availability='sold' WHERE id=?",(ids[0],))
    ref=MarketReferences(db).for_property(p)['groups'][0]
    assert ref['count']==2 and ref['median_sqm'] is None
    assert all(g['count']==0 for g in MarketReferences(db).for_property(p|{'area_basis':'unknown'})['groups'])


@pytest.mark.parametrize('column,value',[('zone','Other'),('area_basis','net'),('currency','USD'),('surface',100),('is_auction',1),('last_seen','2020-01-01T00:00:00+00:00'),('is_demo',1),('availability','unknown')])
def test_references_exclude_incompatible_samples(db,settings,column,value):
    p,ids=populate_comparables(db,settings)
    db.execute(f'UPDATE properties SET {column}=? WHERE id=?',(value,ids[0]))
    ref=MarketReferences(db).for_property(p)['groups'][0]
    assert ref['count']==2 and ref['median_sqm'] is None


def test_workbook_operational_first_with_history_and_safe_criteria(db,settings):
    p,ids=populate_comparables(db,settings)
    p=refresh_analysis(db,p['id'],result(p,prompt='=HYPERLINK("test")'))
    p['screenings']=[{'name':'Ricerca test','fit':True,'custom_prompt':'=HYPERLINK("test")'}]
    blob=export_xlsx([p],db)
    wb=load_workbook(io.BytesIO(blob),data_only=True)
    assert wb.active.title=='Selezione'
    assert {'Selezione','Opportunità','Riferimenti','Comparabili','Criteri AI','Storico','Contatti','Fonti dello stesso asset','Scenari','Metodo e limiti'}==set(wb.sheetnames)
    selection=wb['Selezione'];headers=[cell.value for cell in selection[1]]
    assert selection.cell(2,headers.index('Da ristrutturare /m²')+1).value==2500
    assert wb['Opportunità']['H3'].value==pytest.approx(p['price']/p['surface'])
    assert wb['Criteri AI']['C2'].value.startswith("'=")
    assert wb['Criteri AI']['C2'].data_type=='s'
    assert wb['Storico'].max_row==2 and wb['Comparabili'].max_row==13
    homogeneous=[r for r in wb['Comparabili'].iter_rows(min_row=2,values_only=True) if r[1]=='Stesso stato']
    assert len(homogeneous)==3
    assert selection.auto_filter.ref=='A1:AS2'
    assert load_workbook(io.BytesIO(export_xlsx([]))).active.max_row==1


def test_single_excel_and_detail_reference_contract(api):
    app,c,_=api
    pid=app.state.db.one('SELECT id FROM properties WHERE is_demo=0')['id']
    detail=c.get('/api/properties/'+pid).json()
    assert len(detail['market_references']['groups'])==3
    response=c.get(f'/api/properties/{pid}/memo.xlsx')
    assert response.status_code==200
    wb=load_workbook(io.BytesIO(response.content),data_only=True)
    assert wb.active['A2'].value==pid and wb.active.max_row==2
    assert c.get('/api/properties/not-found/memo.xlsx').status_code==404


def test_truncated_listing_cannot_receive_definitive_custom_verdict(db,settings):
    p=listing(db,settings)|{'description_truncated':True}
    p['analysis']=result(p)
    assert custom_assessment(p,PROMPT)['status']=='uncertain'


def test_references_deduplicate_subject_and_comparables(api):
    app,c,settings=api;db=app.state.db
    p,ids=populate_comparables(db,settings)
    user=c.get('/api/auth/me').json()['user']['id']
    # Agreeing copies count once; disputed copies are covered separately.
    price=db.one('SELECT price FROM properties WHERE id=?',(ids[1],))['price']
    db.execute('UPDATE properties SET price=? WHERE id=?',(price,ids[2]))
    from app.db import now
    for a,b in ((p['id'],ids[0]),(ids[1],ids[2])):
        db.execute('INSERT INTO duplicate_reviews VALUES(?,?,?,?,?) ON CONFLICT(a,b) DO UPDATE SET decision=excluded.decision',(a,b,'same_asset',user,now()))
    group=MarketReferences(db).for_property(p)['groups'][0]
    assert group['count']==1 and group['median_sqm'] is None
    assert group['items'][0]['id'] in ids[1:3]
